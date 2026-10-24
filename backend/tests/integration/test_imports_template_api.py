"""GET /imports/template: the header row and an example, as CSV or XLSX (docs/11 section 4)."""

import csv
import io
from collections.abc import Callable

import pytest
import time_machine
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from salary_api.modules.imports.columns import TEMPLATE_COLUMNS
from tests.factories import make_country
from tests.spreadsheets import XLSX_TYPE, make_import_reference, upload

TEMPLATE = "/api/v1/imports/template"
IMPORTS = "/api/v1/imports"

EXAMPLE = {
    "employee_code": "ACME-010001",
    "first_name": "Jürgen",
    "last_name": "Müller",
    "email": "jurgen.muller@acme.example",
    "job_role": "Software Engineer",
    "level": "L2",
    "employment_type": "FULL_TIME",
    "country": "DE",
    "hire_date": "2022-04-01",
    "currency": "EUR",
    "base_salary": "52000.00",
    "bonus_target": "4160.00",
    "allowance_1_name": "Transport",
    "allowance_1_amount": "1200.00",
    "allowance_2_name": "",
    "allowance_2_amount": "",
    "allowance_3_name": "",
    "allowance_3_amount": "",
    "salary_effective_date": "2022-04-01",
}


def test_the_csv_template_has_the_header_and_one_example_row(
    as_role: Callable[[str], TestClient], db_session: Session
) -> None:
    make_import_reference(db_session)

    response = as_role("HR_MANAGER").get(TEMPLATE, params={"format": "csv"})

    assert response.status_code == 200
    assert response.headers["content-type"] == "text/csv; charset=utf-8"
    assert response.headers["content-disposition"] == (
        'attachment; filename="acme-import-template.csv"'
    )
    assert response.content.startswith("﻿".encode())
    rows = list(csv.reader(io.StringIO(response.content.decode("utf-8-sig"))))
    assert rows == [list(TEMPLATE_COLUMNS), [EXAMPLE[column] for column in TEMPLATE_COLUMNS]]


def test_the_xlsx_template_has_an_employees_sheet_and_a_protected_reference_sheet(
    as_role: Callable[[str], TestClient], db_session: Session
) -> None:
    make_import_reference(db_session)
    make_country(db_session, "GB", is_active=False)

    response = as_role("HR_MANAGER").get(TEMPLATE, params={"format": "xlsx"})

    assert response.headers["content-type"] == XLSX_TYPE
    assert response.headers["content-disposition"] == (
        'attachment; filename="acme-import-template.xlsx"'
    )
    workbook = load_workbook(io.BytesIO(response.content))
    assert workbook.sheetnames == ["Employees", "Reference"]
    employees, reference = workbook["Employees"], workbook["Reference"]
    header = [cell.value for cell in employees[1]]
    assert header == list(TEMPLATE_COLUMNS)
    example = dict(zip(header, (cell.value for cell in employees[2]), strict=True))
    assert (example["employee_code"], example["job_role"], example["country"]) == (
        "ACME-010001",
        "Software Engineer",
        "DE",
    )
    assert reference.protection.sheet is True
    columns = {
        str(column[0].value): [cell.value for cell in column[1:] if cell.value is not None]
        for column in reference.iter_cols()
    }
    assert columns["job_role"] == ["Software Engineer"]
    assert columns["level"] == ["L1", "L2", "L3", "L4", "L5", "L6"]
    assert columns["employment_type"] == ["FULL_TIME", "PART_TIME", "CONTRACTOR"]
    assert columns["country"] == ["DE", "IN", "JP"]  # active countries only
    assert columns["currency"] == ["EUR", "INR", "JPY"]


def test_the_xlsx_template_offers_dropdowns_for_the_enum_columns(
    as_role: Callable[[str], TestClient], db_session: Session
) -> None:
    make_import_reference(db_session)

    content = as_role("HR_MANAGER").get(TEMPLATE, params={"format": "xlsx"}).content

    employees = load_workbook(io.BytesIO(content))["Employees"]
    lists = {
        str(validation.sqref).split(":")[0]: validation.formula1
        for validation in employees.data_validations.dataValidation
        if validation.type == "list"
    }
    assert set(lists) == {"E2", "F2", "G2", "H2", "J2"}  # job role, level, type, country, currency
    assert all(str(formula).startswith("Reference!") for formula in lists.values())


@pytest.mark.parametrize(("file_format", "filename"), [("csv", "t.csv"), ("xlsx", "t.xlsx")])
def test_a_template_uploaded_unchanged_is_a_dry_run_without_errors(
    as_role: Callable[[str], TestClient],
    db_session: Session,
    frozen_time: time_machine.Traveller,
    file_format: str,
    filename: str,
) -> None:
    make_import_reference(db_session)
    hr = as_role("HR_MANAGER")
    template = hr.get(TEMPLATE, params={"format": file_format}).content

    body = hr.post(IMPORTS, **upload(filename, template)).json()

    assert (body["valid_count"], body["errors"]) == (1, [])


def test_an_unknown_template_format_is_422(as_role: Callable[[str], TestClient]) -> None:
    response = as_role("HR_MANAGER").get(TEMPLATE, params={"format": "pdf"})

    assert response.status_code == 422
