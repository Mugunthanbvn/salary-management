"""GET /employees/export: the filtered list as CSV or XLSX, role-aware and formula-safe
(docs/07-api-spec.md section 6.8, docs/11-import-export.md section 6)."""

import csv
import io
from collections.abc import Callable
from datetime import date, datetime
from decimal import Decimal

import pytest
import time_machine
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from salary_api.core.config import Settings
from salary_api.modules.audit.models import AuditEvent
from tests.conftest import signed_in_client
from tests.factories import make_employee, make_fx_rates, make_job_role, make_user
from tests.integration.test_salary_masking import SALARY_QUERIES

EXPORT = "/api/v1/employees/export"
PROBLEMS = "https://acme.example/problems/"
XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
PROFILE_COLUMNS = [
    "employee_code",
    "first_name",
    "last_name",
    "email",
    "job_role",
    "level",
    "employment_type",
    "country",
    "hire_date",
    "department",
    "status",
]
SALARY_COLUMNS = [
    "currency",
    "base_salary",
    "bonus_target",
    "allowances_total",
    "total_comp",
    "base_salary_usd",
    "total_comp_usd",
    "fx_as_of",
]


def _csv_rows(content: bytes) -> list[list[str]]:
    assert content.startswith("﻿".encode())  # so Excel reads it as UTF-8
    return list(csv.reader(io.StringIO(content.decode("utf-8-sig"))))


def _lena(session: Session) -> None:
    make_fx_rates(session)
    make_employee(
        session,
        employee_code="ACME-000001",
        first_name="Lena",
        last_name="Müller",
        email="lena@acme.example",
        job_role=make_job_role(session, "Software Engineer"),
        level="L3",
        hire_date=date(2024, 1, 15),
        base_minor=6_000_000,
        components=[
            ("ALLOWANCE", "Transport", 60_000),
            ("ALLOWANCE", "Meals", 40_000),
            ("BONUS_TARGET", "Target bonus", 600_000),
        ],
    )


def test_hr_exports_a_csv_with_salary_columns_in_major_units(
    as_role: Callable[[str], TestClient], db_session: Session, frozen_time: time_machine.Traveller
) -> None:
    _lena(db_session)

    response = as_role("HR_MANAGER").get(EXPORT, params={"format": "csv"})

    assert response.status_code == 200
    assert response.headers["content-type"] == "text/csv; charset=utf-8"
    assert response.headers["content-disposition"] == (
        'attachment; filename="acme-employees-2026-10-01.csv"'
    )
    assert _csv_rows(response.content) == [
        PROFILE_COLUMNS + SALARY_COLUMNS,
        [
            "ACME-000001",
            "Lena",
            "Müller",
            "lena@acme.example",
            "Software Engineer",
            "L3",
            "FULL_TIME",
            "DE",
            "2024-01-15",
            "Engineering",
            "ACTIVE",
            "EUR",
            "60000.00",
            "6000.00",
            "1000.00",
            "67000.00",
            "65100.00",  # 60,000.00 EUR at 1.085
            "72695.00",
            "2026-10-01",
        ],
    ]


def test_yen_amounts_are_exported_without_decimals(
    as_role: Callable[[str], TestClient], db_session: Session, frozen_time: time_machine.Traveller
) -> None:
    make_fx_rates(db_session)
    make_employee(db_session, country="JP", base_minor=4_500_000)

    [_, row] = _csv_rows(as_role("HR_MANAGER").get(EXPORT).content)

    assert row[11:15] == ["JPY", "4500000", "0", "0"]


def test_a_viewer_exports_no_salary_columns_at_all(
    as_role: Callable[[str], TestClient], db_session: Session, frozen_time: time_machine.Traveller
) -> None:
    _lena(db_session)

    response = as_role("VIEWER").get(EXPORT, params={"format": "csv"})

    assert response.status_code == 200
    rows = _csv_rows(response.content)
    assert rows[0] == PROFILE_COLUMNS
    assert len(rows[1]) == len(PROFILE_COLUMNS)
    assert "60000" not in response.text
    assert "EUR" not in response.text


@pytest.mark.parametrize("params", SALARY_QUERIES)
def test_a_viewer_cannot_export_by_a_salary_filter_or_sort(
    as_role: Callable[[str], TestClient], params: dict[str, str | int]
) -> None:
    response = as_role("VIEWER").get(EXPORT, params=params | {"currency": "EUR"})

    assert response.status_code == 403
    assert response.json()["type"] == f"{PROBLEMS}forbidden"


def test_the_export_applies_the_list_filters_and_sort_without_paging(
    as_role: Callable[[str], TestClient], db_session: Session, frozen_time: time_machine.Traveller
) -> None:
    for day in range(1, 31):
        make_employee(db_session, country="DE", hire_date=date(2024, 1, day))
    make_employee(db_session, country="JP")
    make_employee(db_session, country="DE", status="TERMINATED")

    response = as_role("HR_MANAGER").get(EXPORT, params={"country": "DE", "sort": "-hire_date"})

    hire_dates = [row[8] for row in _csv_rows(response.content)[1:]]
    assert len(hire_dates) == 30  # more than a page; the JP and terminated ones are left out
    assert hire_dates == sorted(hire_dates, reverse=True)


def test_the_export_refuses_paging_and_an_unknown_format(
    as_role: Callable[[str], TestClient],
) -> None:
    hr = as_role("HR_MANAGER")

    assert hr.get(EXPORT, params={"page": 2}).status_code == 422
    assert hr.get(EXPORT, params={"format": "pdf"}).status_code == 422


def test_names_that_look_like_formulas_are_neutralised_in_the_csv(
    as_role: Callable[[str], TestClient], db_session: Session, frozen_time: time_machine.Traveller
) -> None:
    make_employee(db_session, first_name="=1+1", last_name="@SUM(A1)")

    [_, row] = _csv_rows(as_role("VIEWER").get(EXPORT).content)

    assert (row[1], row[2]) == ("'=1+1", "'@SUM(A1)")


def test_hr_exports_an_xlsx_with_typed_and_neutralised_cells(
    as_role: Callable[[str], TestClient], db_session: Session, frozen_time: time_machine.Traveller
) -> None:
    _lena(db_session)
    make_employee(
        db_session, employee_code="ACME-000002", first_name="=HYPERLINK(1)", last_name="Zimmer"
    )  # after Müller in the default sort

    response = as_role("HR_MANAGER").get(EXPORT, params={"format": "xlsx"})

    assert response.headers["content-type"] == XLSX_TYPE
    assert response.headers["content-disposition"] == (
        'attachment; filename="acme-employees-2026-10-01.xlsx"'
    )
    sheet = load_workbook(io.BytesIO(response.content))["Employees"]
    rows = list(sheet.iter_rows())
    assert [cell.value for cell in rows[0]] == PROFILE_COLUMNS + SALARY_COLUMNS
    lena = {
        name: cell for name, cell in zip(PROFILE_COLUMNS + SALARY_COLUMNS, rows[1], strict=True)
    }
    assert (lena["first_name"].value, lena["first_name"].data_type) == ("Lena", "s")
    assert lena["hire_date"].value == datetime(2024, 1, 15)  # a date cell
    assert Decimal(str(lena["base_salary"].value)) == Decimal("60000.00")
    assert lena["base_salary"].data_type == "n"
    other = {
        name: cell for name, cell in zip(PROFILE_COLUMNS + SALARY_COLUMNS, rows[2], strict=True)
    }
    assert (other["first_name"].value, other["first_name"].data_type) == ("'=HYPERLINK(1)", "s")


def test_a_viewer_xlsx_export_has_no_salary_columns(
    as_role: Callable[[str], TestClient], db_session: Session, frozen_time: time_machine.Traveller
) -> None:
    _lena(db_session)

    content = as_role("VIEWER").get(EXPORT, params={"format": "xlsx"}).content

    sheet = load_workbook(io.BytesIO(content))["Employees"]
    assert [cell.value for cell in next(sheet.iter_rows())] == PROFILE_COLUMNS


def test_each_export_writes_an_export_download_event(
    app: FastAPI,
    client: TestClient,
    db_session: Session,
    settings: Settings,
    frozen_time: time_machine.Traveller,
) -> None:
    _lena(db_session)
    viewer = make_user(db_session, role="VIEWER")

    signed_in_client(app, viewer, settings).get(
        EXPORT, params={"format": "xlsx", "country": ["DE", "FR"], "q": "mül"}
    )

    event = db_session.scalars(select(AuditEvent)).one()
    assert (event.action, event.actor_id, event.entity, event.entity_id) == (
        "export.download",
        viewer.id,
        "employee",
        None,
    )
    assert event.diff == {
        "format": "xlsx",
        "filters": {"q": "mül", "country": ["DE", "FR"], "status": ["ACTIVE"]},
        "row_count": 1,
        "salary_included": False,
    }
