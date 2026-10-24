"""Parse an uploaded CSV or XLSX file into raw rows (docs/11-import-export.md section 5.1).

Parsing is the only format-specific step: both parsers return the same `RawRow`s, so the
validator is written and tested once (ADR 0012). Pure: no database, no logging.
"""

import csv
import io
import re
import warnings
import zipfile
from collections.abc import Iterable, Iterator, Mapping, Sequence
from dataclasses import dataclass
from datetime import date, datetime, time
from decimal import Decimal
from typing import TYPE_CHECKING, Final

from salary_api.modules.imports.columns import REQUIRED_COLUMNS, TEMPLATE_COLUMNS

if TYPE_CHECKING:
    from openpyxl.workbook.workbook import Workbook as ReadOnlyWorkbook

type Cell = str | None

NOT_UTF8 = "The file is not UTF-8. In Excel choose File > Save As > 'CSV UTF-8 (Comma delimited)'."
NOT_A_WORKBOOK = (
    "The file is not an Excel workbook (.xlsx) that can be read. Save it again as .xlsx."
)

EMPLOYEES_SHEET: Final = "Employees"
XLSX_COLUMNS_READ: Final = 50
"""Only the first 50 columns of a sheet are read (the template has 19)."""

# <f> or <x:f ...>: a formula element in a worksheet's XML.
_FORMULA_ELEMENT = re.compile(rb"<(?:[A-Za-z_][\w.-]*:)?f[\s/>]")


class UnreadableFile(ValueError):
    """The upload cannot be read as an import file (400 bad-request). The message says why."""


class HeaderMismatch(UnreadableFile):
    """The header row does not match the template: unknown, missing or repeated columns."""


class TooManyRows(ValueError):
    """The file has more data rows than `limit`."""

    def __init__(self, limit: int) -> None:
        super().__init__(f"The file has more than {limit} employee rows.")
        self.limit = limit


@dataclass(frozen=True, slots=True)
class RawRow:
    """One data row: its spreadsheet row number (the header is row 1) and a text value, or
    None, for every template column. Values are not trimmed yet."""

    row: int
    values: Mapping[str, Cell]


def parse_csv(content: bytes, *, max_rows: int) -> list[RawRow]:
    """Rows of a UTF-8 CSV file (a BOM is allowed), comma-separated, quoted as RFC 4180."""
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        raise UnreadableFile(NOT_UTF8) from None
    reader = csv.reader(io.StringIO(text, newline=""), strict=True)
    return _table(_csv_records(reader), max_rows=max_rows)


def parse_xlsx(content: bytes, *, max_rows: int, max_uncompressed_bytes: int) -> list[RawRow]:
    """Rows of the sheet called `Employees`, or else of the first sheet, of an XLSX workbook.

    Values only: openpyxl reads the value Excel saved for each formula and never runs
    anything (`read_only=True, data_only=True`). Refused unread when the archive would
    expand past `max_uncompressed_bytes` (a zip bomb).
    """
    has_formulas = _check_archive(content, max_uncompressed_bytes)
    values = _open_workbook(content, data_only=True)
    # Only to find formulas whose value was never saved, which data_only reads as empty.
    formulas = _open_workbook(content, data_only=False) if has_formulas else None
    try:
        return _table(_xlsx_records(values, formulas), max_rows=max_rows)
    finally:
        values.close()
        if formulas is not None:
            formulas.close()


def _check_archive(content: bytes, max_uncompressed_bytes: int) -> bool:
    """Whether a worksheet has a formula. Raises UnreadableFile for a broken or huge zip."""
    try:
        archive = zipfile.ZipFile(io.BytesIO(content))
    except (zipfile.BadZipFile, OSError, EOFError, ValueError):
        raise UnreadableFile(NOT_A_WORKBOOK) from None
    with archive:
        entries = archive.infolist()
        # The declared sizes: zipfile never inflates an entry past its declared size.
        if sum(entry.file_size for entry in entries) > max_uncompressed_bytes:
            raise UnreadableFile(
                f"The workbook expands to more than {_size(max_uncompressed_bytes)} "
                "when unzipped. Save only the employee rows in a new workbook."
            )
        try:
            return any(
                _FORMULA_ELEMENT.search(archive.read(entry)) is not None
                for entry in entries
                if entry.filename.startswith("xl/worksheets/")
            )
        except Exception:  # a damaged entry: bad CRC, broken deflate stream, ...
            raise UnreadableFile(NOT_A_WORKBOOK) from None


def _open_workbook(content: bytes, *, data_only: bool) -> "ReadOnlyWorkbook":
    # Imported here, not at the top: openpyxl is slow to import (docs/19 section 6.3).
    from openpyxl import load_workbook  # noqa: PLC0415

    try:
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")  # e.g. "unsupported extension": not the user's error
            return load_workbook(io.BytesIO(content), read_only=True, data_only=data_only)
    except Exception:  # openpyxl raises many kinds of errors for a damaged file
        raise UnreadableFile(NOT_A_WORKBOOK) from None


def _xlsx_records(
    values: "ReadOnlyWorkbook", formulas: "ReadOnlyWorkbook | None"
) -> Iterator[tuple[int, Sequence[Cell]]]:
    """Numbered rows of text cells, read side by side from the values and formulas workbooks."""
    value_rows = _sheet_rows(values)
    formula_rows = _sheet_rows(formulas) if formulas is not None else None
    number = 0
    while True:
        number += 1
        try:
            cells = next(value_rows, None)
            formula_cells = next(formula_rows, None) if formula_rows is not None else None
        except Exception:  # a damaged sheet only shows while it is read
            raise UnreadableFile(NOT_A_WORKBOOK) from None
        if cells is None:
            return
        for position, (value, formula) in enumerate(zip(cells, formula_cells or (), strict=False)):
            if value is None and isinstance(formula, str) and formula.startswith("="):
                raise UnreadableFile(
                    f"Cell {_letter(position)}{number} has a formula without a saved value. "
                    "Copy the cells and paste them as values, then upload again."
                )
        yield number, [_cell_text(value) for value in cells]


def _sheet_rows(workbook: "ReadOnlyWorkbook") -> Iterator[tuple[object, ...]]:
    try:
        names: list[str] = workbook.sheetnames
        sheet = workbook[EMPLOYEES_SHEET] if EMPLOYEES_SHEET in names else workbook.worksheets[0]
        # The saved dimensions can be wrong; read the rows that are really there.
        sheet.reset_dimensions()
        return sheet.iter_rows(max_col=XLSX_COLUMNS_READ, values_only=True)
    except Exception:
        raise UnreadableFile(NOT_A_WORKBOOK) from None


def _cell_text(value: object) -> Cell:
    """A cell's value as the text a CSV of the same sheet holds."""
    if value is None or isinstance(value, str):
        return value
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float):
        # Through the float's shortest repr: 1234.56 stays "1234.56", never 1234.5599999...
        return format(Decimal(repr(value)), "f")
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == time() else value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value)


def _size(byte_count: int) -> str:
    megabytes = byte_count // (1024 * 1024)
    return f"{megabytes} MB" if megabytes else f"{byte_count} bytes"


def _csv_records(reader: Iterator[list[str]]) -> Iterator[tuple[int, Sequence[Cell]]]:
    """Numbered records. A multi-line quoted value stays one row, as Excel shows it."""
    number = 0
    while True:
        number += 1
        try:
            cells = next(reader)
        except StopIteration:
            return
        except csv.Error as error:
            raise UnreadableFile(f"Row {number} could not be read: {error}.") from None
        yield number, cells


def _table(records: Iterable[tuple[int, Sequence[Cell]]], *, max_rows: int) -> list[RawRow]:
    """The data rows under the header row, keyed by template column; blank rows skipped."""
    numbered = iter(records)
    first = next(numbered, None)
    if first is None:
        raise UnreadableFile("The file is empty: it needs a header row and employee rows.")
    columns = _columns(first[1])
    rows: list[RawRow] = []
    for number, cells in numbered:
        values = _row_values(number, columns, cells)
        if values is None:
            continue
        if len(rows) == max_rows:
            raise TooManyRows(max_rows)
        rows.append(RawRow(row=number, values=values))
    if not rows:
        raise UnreadableFile("The file has a header row but no employee rows.")
    return rows


def _columns(header: Sequence[Cell]) -> list[str]:
    """The template column of each position ("" for a column with no header).

    Raises HeaderMismatch naming every unknown, repeated and missing column.
    """
    columns = [(cell or "").strip().lower() for cell in header]
    named = [column for column in columns if column]
    problems: list[str] = []
    if len(named) == 1 and ";" in named[0]:
        problems.append("Separate the columns with commas, not semicolons.")
    else:
        unknown = [column for column in named if column not in TEMPLATE_COLUMNS]
        if unknown:
            problems.append(f"Unknown column{_s(unknown)}: {', '.join(unknown)}.")
        repeated = sorted({column for column in named if named.count(column) > 1})
        problems += [f"The column {column} appears more than once." for column in repeated]
    missing = [column for column in TEMPLATE_COLUMNS if column in REQUIRED_COLUMNS - set(named)]
    if missing:
        problems.append(f"Missing required column{_s(missing)}: {', '.join(missing)}.")
    if problems:
        raise HeaderMismatch(" ".join(problems))
    return columns


def _row_values(
    number: int, columns: Sequence[str], cells: Sequence[Cell]
) -> dict[str, Cell] | None:
    """The row's value for every template column, or None for a blank row."""
    values: dict[str, Cell] = dict.fromkeys(TEMPLATE_COLUMNS)
    blank = True
    for position, cell in enumerate(cells):
        value = _value(cell)
        if value is None:
            continue
        column = columns[position] if position < len(columns) else ""
        if not column:
            raise UnreadableFile(
                f"Row {number} has a value in column {_letter(position)}, which has no header. "
                "Put values that contain commas in double quotes."
            )
        values[column] = value
        blank = False
    return None if blank else values


def _value(cell: Cell) -> Cell:
    """None for an empty or blank cell, else the text as written."""
    return cell if cell is not None and cell.strip() else None


def _letter(position: int) -> str:
    """The spreadsheet column letter of a 0-based position: 0 is A, 26 is AA."""
    letters = ""
    number = position + 1
    while number:
        number, remainder = divmod(number - 1, 26)
        letters = chr(ord("A") + remainder) + letters
    return letters


def _s(items: Sequence[object]) -> str:
    return "s" if len(items) > 1 else ""
